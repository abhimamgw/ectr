"""playground streamlit app."""

import pandas as pd
from openpyxl import load_workbook
import streamlit as st

st.set_page_config(
    page_title="eCTR Conversion Tool",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.title("eCTR Conversion Tool")
st.write("Extracts hours for Rigid Pipeline and Rigid Spool")


def get_activity_data(file_path, sheet_name):
    if not sheet_name.startswith("Rigid"):
        return pd.DataFrame()

    df = pd.read_excel(file_path, sheet_name=sheet_name)
    if not df.iloc[:, 0].eq("Activities").any():
        return pd.DataFrame()

    idx_start = df.index[df.iloc[:, 0] == "Activities"][0] + 2
    idx_end = df.index[df.iloc[:, 0] == "Inputs"][0] - 2
    hours = df[idx_start:idx_end]
    hours.columns = df.iloc[idx_start - 1, :]
    hours = hours[hours["Activity WBS"].isna()]
    hours = hours[["Block", "Scope", str("Engineering hours\nTotal")]]
    hours.columns = ["Block", "Scope", "Hours"]
    hours = hours[hours["Hours"] > 0]
    hours = hours.groupby(["Block", "Scope"])["Hours"].sum().reset_index()
    hours.insert(0, "CTR", sheet_name)
    return hours


def main(file_path):
    with st.spinner("Running..."):
        sheets = load_workbook(filename=file_path, read_only=True).sheetnames
        data = pd.concat([get_activity_data(file_path, sheet) for sheet in sheets])
        data = data.reset_index(drop=True)
        data["Hours [%]"] = data["Hours"] / data["Hours"].sum() * 100
        st.dataframe(data, hide_index=True, use_container_width=True)


if __name__ == "__main__":
    # file_path = "eCTR_Katlan.xlsx"
    file_path = st.sidebar.file_uploader("Upload eCTR file", type="xlsx")
    if file_path:
        main(file_path)